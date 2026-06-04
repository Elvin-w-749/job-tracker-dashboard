from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Query, Header, Response
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from sqlalchemy import func, inspect as sa_inspect, text as sa_text
from database import engine, get_db, Base, SessionLocal
from database import migrate_add_user_id
from models import Application, UserConfig, BatchTask, BatchResult
from schemas import ApplicationCreate, APIKeyRequest
from services.minimax_service import MiniMaxService
from services.file_parser_service import parse_file_by_extension
from datetime import date, datetime, timedelta
import json
import httpx
import uuid
import asyncio
import csv
import io
import urllib.parse
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

ALLOWED_SORT_FIELDS = {"apply_date", "company_name", "match_score", "created_at", "position", "city"}

Base.metadata.create_all(bind=engine)
migrate_add_user_id()

# 数据库迁移：检查 company_tier 列是否存在，不存在则添加
_inspector = sa_inspect(engine)
_columns = [col["name"] for col in _inspector.get_columns("applications")]
if "company_tier" not in _columns:
    with engine.connect() as conn:
        conn.execute(
            sa_text(
                "ALTER TABLE applications ADD COLUMN company_tier VARCHAR"
            )
        )
        conn.commit()


def extract_company_tier(company_size: str) -> str:
    """从企业规模字符串中提取企业档次分类"""
    if not company_size:
        return "未知"
    if "中小厂" in company_size or "100-499" in company_size:
        return "中小厂"
    if "中大厂" in company_size or "2000-4999" in company_size:
        return "中大厂"
    if "小厂" in company_size or "100人" in company_size or "少于" in company_size:
        return "小厂"
    if "中厂" in company_size or "500-1999" in company_size:
        return "中厂"
    if "大厂" in company_size or "10000人以上" in company_size:
        return "大厂"
    return "未知"

app = FastAPI(title="Job Tracker API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://localhost:3001"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_user_id(x_user_id: str = Header(default="default")) -> str:
    return x_user_id or "default"


def get_api_key(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "minimax_api_key").first()
    if not config or not config.value:
        raise HTTPException(status_code=401, detail="API Key 未配置，请在设置中配置 MiniMax API Key")
    return config.value

def get_minimax_model(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)) -> str | None:
    config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "minimax_model_name").first()
    return config.value if config else None


def get_resume_profile(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)) -> str | None:
    config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "resume_profile").first()
    return config.value if config else None


@app.get("/api/health")
async def health_check():
    return {"status": "ok"}


@app.post("/api/settings/api-key")
async def save_api_key(request: APIKeyRequest, db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "minimax_api_key").first()
    if config:
        config.value = request.api_key
    else:
        config = UserConfig(user_id=user_id, key="minimax_api_key", value=request.api_key)
        db.add(config)
        
    if request.model_name:
        model_config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "minimax_model_name").first()
        if model_config:
            model_config.value = request.model_name
        else:
            model_config = UserConfig(user_id=user_id, key="minimax_model_name", value=request.model_name)
            db.add(model_config)

    db.commit()
    return {"success": True, "message": "配置保存成功"}


@app.post("/api/settings/api-key/clear")
async def clear_api_key_post(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "minimax_api_key").delete()
    db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "minimax_model_name").delete()
    db.commit()
    return {"success": True, "message": "配置已重置"}


@app.get("/api/settings/api-key")
async def get_api_key_status(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "minimax_api_key").first()
    model_config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "minimax_model_name").first()
    model_name = model_config.value if model_config else "MiniMax-Text-01"

    if config and config.value:
        preview = config.value[-8:] if len(config.value) > 8 else config.value
        return {"success": True, "data": {"configured": True, "key_preview": f"***{preview}", "model_name": model_name}}
    return {"success": True, "data": {"configured": False, "key_preview": None, "model_name": model_name}}


@app.post("/api/settings/api-key/test")
async def test_api_key(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "minimax_api_key").first()
    if not config or not config.value:
        return {"success": False, "message": "API Key 未配置"}
    try:
        async with httpx.AsyncClient(timeout=10.0) as client:
            headers = {"Authorization": f"Bearer {config.value}"}
            response = await client.get("https://api.minimaxi.com/v1/models", headers=headers)
            if response.status_code == 200:
                return {"success": True, "message": "连接成功"}
            return {"success": False, "message": "API Key 无效"}
    except Exception as e:
        return {"success": False, "message": f"测试失败: {str(e)}"}


@app.post("/api/settings/resume-profile")
async def save_resume_profile(request: dict, db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    profile = request.get("resume_profile", "").strip()
    config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "resume_profile").first()
    if config:
        config.value = profile
    else:
        config = UserConfig(user_id=user_id, key="resume_profile", value=profile)
        db.add(config)
    db.commit()
    return {"success": True, "message": "简历背景保存成功"}


@app.post("/api/settings/clear-data")
async def clear_all_data(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    """清除当前用户数据（投递记录、批量任务、简历背景），保留 API Key"""
    db.query(Application).filter(Application.user_id == user_id).delete()
    db.query(BatchResult).filter(BatchResult.user_id == user_id).delete()
    db.query(BatchTask).filter(BatchTask.user_id == user_id).delete()
    resume_config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "resume_profile").first()
    if resume_config:
        resume_config.value = ""
    db.commit()
    return {"success": True, "message": "当前用户数据已清除，API Key 已保留"}


@app.get("/api/settings/resume-profile")
async def get_resume_profile_status(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "resume_profile").first()
    return {
        "success": True,
        "data": {
            "configured": bool(config and config.value),
            "profile": config.value if config else None,
        }
    }


@app.post("/api/resume/analyze")
async def analyze_resume(
    file: UploadFile = File(...),
    api_key: str = Depends(get_api_key),
    user_id: str = Depends(get_user_id)
):
    """上传简历文件（图片/PDF/DOCX/TXT），AI 提取求职者背景摘要"""
    filename = file.filename or ""
    data = await file.read()
    if not data:
        return {"success": False, "error": "文件内容为空"}

    ext = filename.lower().split(".")[-1] if "." in filename else ""
    service = MiniMaxService(api_key)

    try:
        if ext in ("png", "jpg", "jpeg", "webp", "gif", "bmp"):
            profile_text = await service.analyze_resume_image(data, file.content_type)
            return {"success": True, "data": {"profile": profile_text.strip()}}

        if ext in ("docx",):
            from services.file_parser_service import parse_docx
            resume_text = parse_docx(data)
            if not resume_text.strip():
                return {"success": False, "error": "未能从文件中提取到文本内容"}
            profile_text = await service.analyze_resume_text(resume_text)
            return {"success": True, "data": {"profile": profile_text.strip()}}

        if ext in ("txt", "md", ""):
            from services.file_parser_service import parse_text
            resume_text = parse_text(data)
            if not resume_text.strip():
                return {"success": False, "error": "未能从文件中提取到文本内容"}
            profile_text = await service.analyze_resume_text(resume_text)
            return {"success": True, "data": {"profile": profile_text.strip()}}

        if ext in ("pdf",):
            return {"success": False, "error": "PDF 简历暂不支持，请先将 PDF 转为图片（JPG/PNG）后上传"}

        return {"success": False, "error": f"不支持的简历格式: .{ext}，请上传图片( JPG/PNG)或文档(DOCX/TXT/MD)"}
    except Exception as e:
        return {"success": False, "error": f"简历解析失败: {str(e)}"}


@app.post("/api/resume/confirm-profile")
async def confirm_resume_profile(request: dict, db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    """确认保存 AI 解析的简历背景到 user_config"""
    profile = request.get("profile", "").strip()
    if not profile:
        return {"success": False, "error": "简历背景不能为空"}

    config = db.query(UserConfig).filter(UserConfig.user_id == user_id, UserConfig.key == "resume_profile").first()
    if config:
        config.value = profile
    else:
        config = UserConfig(user_id=user_id, key="resume_profile", value=profile)
        db.add(config)
    db.commit()
    return {"success": True, "message": "简历背景已保存，后续 JD 分析将基于此背景评估匹配度"}


@app.post("/api/jd/analyze")
async def analyze_jd(
    image: UploadFile = File(...),
    api_key: str = Depends(get_api_key),
    resume_profile: str | None = Depends(get_resume_profile),
    model_name: str | None = Depends(get_minimax_model),
    user_id: str = Depends(get_user_id)
):
    if not image.content_type or not image.content_type.startswith("image/"):
        return {"success": False, "error": "请上传图片文件"}

    try:
        image_data = await image.read()
        service = MiniMaxService(api_key=api_key, resume_profile=resume_profile, model_name=model_name)
        result = await service.analyze_jd_image(image_data, image.content_type)

        return {
            "success": True,
            "data": result
        }
    except Exception as e:
        return {"success": False, "error": f"图片解析失败: {str(e)}"}


@app.get("/api/applications")
async def get_applications(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=5000),
    sort_by: str = Query("apply_date"),
    sort_order: str = Query("desc"),
    company_name: str = Query(None),
    position: str = Query(None),
    city: str = Query(None),
    company_size: str = Query(None),
    status: str = Query(None),
    channel: str = Query(None),
    date_from: date = Query(None),
    date_to: date = Query(None),
    match_score_min: int = Query(None),
    match_score_max: int = Query(None),
    keyword: str = Query(None),
    db: Session = Depends(get_db),
    user_id: str = Depends(get_user_id)
):
    if sort_by not in ALLOWED_SORT_FIELDS:
        raise HTTPException(status_code=400, detail=f"不支持的排序字段: {sort_by}，可选值: {', '.join(ALLOWED_SORT_FIELDS)}")

    query = db.query(Application).filter(Application.user_id == user_id)

    if company_name:
        query = query.filter(Application.company_name.contains(company_name))
    if position:
        query = query.filter(Application.position.contains(position))
    if city:
        query = query.filter(Application.city == city)
    if company_size:
        query = query.filter(Application.company_size == company_size)
    if status:
        query = query.filter(Application.status == status)
    if channel:
        query = query.filter(Application.channel == channel)
    if date_from:
        query = query.filter(Application.apply_date >= date_from)
    if date_to:
        query = query.filter(Application.apply_date <= date_to)
    if match_score_min is not None:
        query = query.filter(Application.match_score >= match_score_min)
    if match_score_max is not None:
        query = query.filter(Application.match_score <= match_score_max)
    if keyword:
        query = query.filter(Application.jd_text.contains(keyword))

    total = query.count()

    sort_column = getattr(Application, sort_by)
    if sort_order == "desc":
        query = query.order_by(sort_column.desc())
    else:
        query = query.order_by(sort_column.asc())

    items = query.offset((page - 1) * page_size).limit(page_size).all()

    return {
        "success": True,
        "data": {
            "items": [
                {
                    "id": item.id,
                    "company_name": item.company_name,
                    "position": item.position,
                    "city": item.city,
                    "company_size": item.company_size,
                    "company_tier": item.company_tier,
                    "jd_text": item.jd_text,
                    "apply_date": item.apply_date.isoformat() if item.apply_date else None,
                    "status": item.status,
                    "channel": item.channel,
                    "salary_range": item.salary_range,
                    "education": item.education,
                    "experience": item.experience,
                    "skills": json.loads(item.skills) if item.skills else None,
                    "match_score": item.match_score,
                    "created_at": item.created_at.isoformat() if item.created_at else None,
                }
                for item in items
            ],
            "total": total,
            "page": page,
            "page_size": page_size
        }
    }


@app.get("/api/stats/overview")
async def get_stats_overview(
    company_name: str = Query(None),
    position: str = Query(None),
    city: str = Query(None),
    status: str = Query(None),
    channel: str = Query(None),
    match_score_min: int = Query(None),
    match_score_max: int = Query(None),
    db: Session = Depends(get_db),
    user_id: str = Depends(get_user_id)
):
    query = db.query(Application).filter(Application.user_id == user_id)

    if company_name:
        query = query.filter(Application.company_name.contains(company_name))
    if position:
        query = query.filter(Application.position.contains(position))
    if city:
        query = query.filter(Application.city == city)
    if status:
        query = query.filter(Application.status == status)
    if channel:
        query = query.filter(Application.channel == channel)
    if match_score_min is not None:
        query = query.filter(Application.match_score >= match_score_min)
    if match_score_max is not None:
        query = query.filter(Application.match_score <= match_score_max)

    total = query.count()

    city_results = query.with_entities(Application.city, func.count(Application.id)).group_by(Application.city).all()
    city_distribution = {c or "未知": cnt for c, cnt in city_results}

    status_results = query.with_entities(Application.status, func.count(Application.id)).group_by(Application.status).all()
    status_distribution = {s or "未知": cnt for s, cnt in status_results}

    size_results = query.with_entities(Application.company_size, func.count(Application.id)).group_by(Application.company_size).all()
    company_size_distribution = {sz or "未知": cnt for sz, cnt in size_results}

    tier_results = query.with_entities(Application.company_tier, func.count(Application.id)).group_by(Application.company_tier).all()
    company_tier_distribution = {t or "未知": cnt for t, cnt in tier_results}

    channel_results = query.with_entities(Application.channel, func.count(Application.id)).group_by(Application.channel).all()
    channel_distribution = {ch or "未知": cnt for ch, cnt in channel_results}

    avg_score = query.with_entities(func.avg(Application.match_score)).scalar() or 0

    match_distribution = {
        "0-20": query.filter(Application.match_score.between(0, 20)).count(),
        "21-40": query.filter(Application.match_score.between(21, 40)).count(),
        "41-60": query.filter(Application.match_score.between(41, 60)).count(),
        "61-80": query.filter(Application.match_score.between(61, 80)).count(),
        "81-100": query.filter(Application.match_score.between(81, 100)).count(),
    }

    timeline_results = query.filter(Application.apply_date.isnot(None)).with_entities(
        Application.apply_date,
        func.count(Application.id)
    ).group_by(Application.apply_date).order_by(Application.apply_date).all()
    timeline = [{"date": str(d), "count": c} for d, c in timeline_results]

    # 覆盖公司数
    companies_covered = query.with_entities(
        func.count(func.distinct(Application.company_name))
    ).scalar() or 0

    # 已回复率 = 非"已投递"且非"7天未回复"且非"已读不回" 的数量占比
    no_response_statuses = ["已投递", "7天未回复", "已读不回"]
    responded_count = query.filter(
        Application.status.notin_(no_response_statuses)
    ).count()
    response_rate = round(responded_count / total * 100, 1) if total > 0 else 0.0

    # 求职转化漏斗数据 (15个状态完整映射)
    # 所有状态："已投递", "简历被查看", "已读不回", "7天未回复", "面试中", "一面", "二面", "三面", "HR面", "已offer", "offer谈判中", "已接受offer", "已拒绝offer", "流程结束", "已拒绝"
    
    view_statuses = [
        "简历被查看", "已读不回", "面试中", "一面", "二面", "三面", "HR面",
        "已offer", "offer谈判中", "已接受offer", "已拒绝offer", "流程结束", "已拒绝"
    ]
    interview_statuses = [
        "面试中", "一面", "二面", "三面", "HR面",
        "已offer", "offer谈判中", "已接受offer", "已拒绝offer"
    ]
    offer_statuses = ["已offer", "offer谈判中", "已接受offer", "已拒绝offer"]
    accept_statuses = ["已接受offer"]

    funnel_data = {
        "applications": total,
        "views": query.filter(Application.status.in_(view_statuses)).count(),
        "interviews": query.filter(Application.status.in_(interview_statuses)).count(),
        "offers": query.filter(Application.status.in_(offer_statuses)).count(),
        "accepts": query.filter(Application.status.in_(accept_statuses)).count(),
    }

    # 月度投递趋势
    monthly_results = query.filter(Application.apply_date.isnot(None)).with_entities(
        func.strftime("%Y-%m", Application.apply_date).label("month"),
        func.count(Application.id)
    ).group_by("month").order_by("month").all()
    monthly_trend = [{"month": m, "count": c} for m, c in monthly_results if m]

    # 匹配度 TOP 10
    top_match_results = query.filter(
        Application.match_score.isnot(None)
    ).order_by(Application.match_score.desc()).limit(10).all()
    top_matches = [
        {
            "id": r.id,
            "company_name": r.company_name,
            "position": r.position,
            "city": r.city or "",
            "match_score": r.match_score,
            "status": r.status,
            "channel": r.channel or "",
            "apply_date": str(r.apply_date) if r.apply_date else "",
            "company_tier": r.company_tier or "",
        }
        for r in top_match_results
    ]

    return {
        "success": True,
        "data": {
            "total": total,
            "city_distribution": city_distribution,
            "status_distribution": status_distribution,
            "company_size_distribution": company_size_distribution,
            "company_tier_distribution": company_tier_distribution,
            "channel_distribution": channel_distribution,
            "match_score_avg": round(float(avg_score), 1),
            "match_score_distribution": match_distribution,
            "timeline": timeline,
            "companies_covered": companies_covered,
            "response_rate": response_rate,
            "funnel_data": funnel_data,
            "monthly_trend": monthly_trend,
            "top_matches": top_matches,
        }
    }


@app.post("/api/jd/batch-upload")
async def batch_upload(
    images: list[UploadFile] = File(...),
    api_key: str = Depends(get_api_key),
    resume_profile: str | None = Depends(get_resume_profile),
    model_name: str | None = Depends(get_minimax_model),
    db: Session = Depends(get_db),
    user_id: str = Depends(get_user_id)
):
    task_id = str(uuid.uuid4())

    task = BatchTask(id=task_id, user_id=user_id, total=len(images), status="processing")
    db.add(task)
    db.commit()

    for i in range(len(images)):
        result = BatchResult(task_id=task_id, user_id=user_id, image_index=i, status="pending")
        db.add(result)
    db.commit()

    image_data_list = []
    for img in images:
        data = await img.read()
        image_data_list.append(data)

    asyncio.create_task(_process_batch(task_id, image_data_list, api_key, resume_profile, user_id, model_name))

    return {
        "success": True,
        "data": {
            "task_id": task_id,
            "total": len(images),
            "completed": 0,
            "failed": 0,
            "status": "processing"
        }
    }


async def _process_batch(task_id: str, image_data_list: list[bytes], api_key: str, resume_profile: str | None, user_id: str, model_name: str | None = None):
    """并发批量处理 JD 图片分析，每批最多 10 张并发"""
    service = MiniMaxService(api_key=api_key, resume_profile=resume_profile, model_name=model_name)
    db = SessionLocal()
    semaphore = asyncio.Semaphore(10)

    async def process_one(i: int, image_data: bytes):
        async with semaphore:
            result = db.query(BatchResult).filter_by(task_id=task_id, user_id=user_id, image_index=i).first()
            if not result:
                return
            result.status = "processing"
            db.commit()
            try:
                analysis = await service.analyze_jd_image(image_data)
                result.status = "completed"
                result.result_data = json.dumps(analysis, ensure_ascii=False)
            except Exception as e:
                result.status = "failed"
                result.error_message = str(e)
            db.commit()

    try:
        await asyncio.gather(*[process_one(i, data) for i, data in enumerate(image_data_list)])
        task = db.query(BatchTask).filter_by(id=task_id, user_id=user_id).first()
        if task:
            task.completed = db.query(BatchResult).filter_by(task_id=task_id, user_id=user_id, status="completed").count()
            task.failed = db.query(BatchResult).filter_by(task_id=task_id, user_id=user_id, status="failed").count()
            task.status = "completed"
            task.completed_at = datetime.utcnow()
            db.commit()
    finally:
        db.close()


@app.get("/api/jd/batch-status/{task_id}")
async def get_batch_status(task_id: str, db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    task = db.query(BatchTask).filter_by(id=task_id, user_id=user_id).first()
    if not task:
        return {"success": False, "error": "任务不存在"}

    results = db.query(BatchResult).filter_by(task_id=task_id, user_id=user_id).order_by(BatchResult.image_index).all()

    return {
        "success": True,
        "data": {
            "task_id": task.id,
            "total": task.total,
            "completed": task.completed,
            "failed": task.failed,
            "status": task.status,
            "results": [
                {
                    "index": r.image_index,
                    "status": r.status,
                    "data": json.loads(r.result_data) if r.result_data else None,
                    "error_message": r.error_message
                }
                for r in results
            ]
        }
    }


@app.post("/api/applications/confirm")
async def confirm_application(data: ApplicationCreate, db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    try:
        skills_str = json.dumps(data.skills, ensure_ascii=False) if data.skills else None

        company_tier = data.company_tier or extract_company_tier(data.company_size)

        app_record = Application(
            user_id=user_id,
            company_name=data.company_name,
            position=data.position,
            city=data.city,
            company_size=data.company_size,
            company_tier=company_tier,
            jd_text=data.jd_text,
            salary_range=data.salary_range,
            education=data.education,
            experience=data.experience,
            skills=skills_str,
            match_score=data.match_score,
            status=data.status or "已投递",
            channel=data.channel,
            apply_date=date.today()
        )
        db.add(app_record)
        db.commit()
        db.refresh(app_record)

        return {
            "success": True,
            "data": {
                "id": app_record.id,
                "company_name": app_record.company_name,
                "position": app_record.position,
            }
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


# 全量投递进度枚举（前后端共用）
ALL_STATUSES = [
    "已投递", "简历被查看", "已读不回", "7天未回复",
    "一面", "二面", "三面", "HR面",
    "已offer", "offer谈判中", "已接受offer", "已拒绝offer",
    "流程结束", "已拒绝", "面试中",
]

# 导入时状态直接透传，不强制映射
_STATUS_MAP: dict[str, str] = {}

# 投递渠道映射
_CHANNEL_MAP = {
    "boss": "BOSS直聘",
    "BOSS": "BOSS直聘",
    "Boss": "BOSS直聘",
    "官网": "官网",
    "内推": "内推",
    "猎聘": "猎聘",
    "拉勾": "拉勾",
    "智联": "智联",
    "前程无忧": "前程无忧",
    "实习僧": "实习僧",
    "脉脉": "脉脉",
}


def _map_status(raw: str) -> str:
    """将CSV中的投递进度映射为标准状态"""
    if not raw:
        return "已投递"
    return raw.strip()


def _map_channel(raw: str) -> str:
    """将CSV中的投递渠道映射为标准渠道名"""
    if not raw:
        return None
    raw = raw.strip()
    if raw in _CHANNEL_MAP:
        return _CHANNEL_MAP[raw]
    return raw


@app.post("/api/import/csv")
async def import_csv(file: UploadFile = File(...), db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    """导入CSV文件到投递记录表"""
    filename = file.filename or ""
    ext = filename.lower().split(".")[-1] if "." in filename else ""
    if ext not in ("csv", "xlsx", "xls"):
        return {"success": False, "error": "不支持的文件格式，请上传 .csv、.xlsx 或 .xls 文件"}

    try:
        content = await file.read()
        # 尝试UTF-8解码，如果失败尝试GBK
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                text = content.decode("gbk")

        reader = csv.DictReader(io.StringIO(text))
        imported = 0
        skipped = 0

        for row in reader:
            company_name = row.get("企业名称", "").strip()
            position = row.get("岗位", "").strip()
            if not company_name or not position:
                skipped += 1
                continue

            city = row.get("城市", "").strip() or None
            company_size = row.get("企业规模", "").strip() or None
            jd_text = row.get("JD", "").strip() or None
            date_str = row.get("日期", "").strip()
            raw_status = row.get("投递进度", "").strip()
            raw_channel = row.get("投递渠道", "").strip()

            # 解析日期
            apply_date = None
            if date_str:
                try:
                    apply_date = datetime.strptime(date_str, "%Y/%m/%d").date()
                except ValueError:
                    try:
                        apply_date = datetime.strptime(date_str, "%Y-%m-%d").date()
                    except ValueError:
                        pass

            status = _map_status(raw_status)
            channel = _map_channel(raw_channel)
            company_tier = extract_company_tier(company_size)

            # 重复检测：同公司+同岗位+同日期视为重复
            existing = db.query(Application).filter(
                Application.user_id == user_id,
                Application.company_name == company_name,
                Application.position == position,
                Application.apply_date == apply_date
            ).first()
            if existing:
                skipped += 1
                continue

            app_record = Application(
                user_id=user_id,
                company_name=company_name,
                position=position,
                city=city,
                company_size=company_size,
                company_tier=company_tier,
                jd_text=jd_text,
                apply_date=apply_date,
                status=status,
                channel=channel,
            )
            db.add(app_record)
            imported += 1

        db.commit()

        return {
            "success": True,
            "data": {
                "imported": imported,
                "skipped": skipped,
            }
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


@app.post("/api/jd/analyze-file")
async def analyze_jd_file(
    file: UploadFile = File(...),
    api_key: str = Depends(get_api_key),
    resume_profile: str | None = Depends(get_resume_profile),
    model_name: str | None = Depends(get_minimax_model),
    user_id: str = Depends(get_user_id)
):
    """多格式文件智能分析入口：支持 .csv/.xlsx/.docx/.txt/.md/.png/.jpg"""
    filename = file.filename or ""
    data = await file.read()

    if not data:
        return {"success": False, "error": "文件内容为空"}

    ext = filename.lower().split(".")[-1] if "." in filename else ""

    if ext in ("png", "jpg", "jpeg", "webp", "gif"):
        service = MiniMaxService(api_key=api_key, resume_profile=resume_profile, model_name=model_name)
        try:
            result = await service.analyze_jd_image(data, file.content_type)
            return {"success": True, "data": result}
        except Exception as e:
            return {"success": False, "error": f"图片解析失败: {str(e)}"}

    if ext in ("csv", "xlsx", "xls"):
        try:
            parsed = parse_file_by_extension(data, filename)
            rows = parsed.get("rows", [])
            imported = 0
            skipped = 0
            db = SessionLocal()
            try:
                for row in rows:
                    company_name = str(row.get("企业名称") or row.get("company_name") or "").strip()
                    position = str(row.get("岗位") or row.get("position") or "").strip()
                    if not company_name or not position:
                        skipped += 1
                        continue

                    city = str(row.get("城市") or row.get("city") or "").strip() or None
                    company_size = str(row.get("企业规模") or row.get("company_size") or "").strip() or None
                    jd_text = str(row.get("JD") or row.get("jd_text") or "").strip() or None
                    date_str = str(row.get("日期") or row.get("apply_date") or "").strip()
                    raw_status = str(row.get("投递进度") or row.get("status") or "").strip()
                    raw_channel = str(row.get("投递渠道") or row.get("channel") or "").strip()

                    apply_date = None
                    if date_str:
                        for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y年%m月%d日"):
                            try:
                                apply_date = datetime.strptime(date_str, fmt).date()
                                break
                            except ValueError:
                                continue

                    status = _map_status(raw_status)
                    channel = _map_channel(raw_channel)
                    company_tier = extract_company_tier(company_size)

                    existing = db.query(Application).filter(
                        Application.user_id == user_id,
                        Application.company_name == company_name,
                        Application.position == position,
                        Application.apply_date == apply_date
                    ).first()
                    if existing:
                        skipped += 1
                        continue

                    app_record = Application(
                        user_id=user_id,
                        company_name=company_name,
                        position=position,
                        city=city,
                        company_size=company_size,
                        company_tier=company_tier,
                        jd_text=jd_text,
                        apply_date=apply_date,
                        status=status,
                        channel=channel,
                    )
                    db.add(app_record)
                    imported += 1

                db.commit()
            finally:
                db.close()

            return {
                "success": True,
                "data": {"imported": imported, "skipped": skipped, "type": "structured"}
            }
        except Exception as e:
            return {"success": False, "error": f"表格解析失败: {str(e)}"}

    if ext in ("docx", "txt", "md"):
        try:
            parsed = parse_file_by_extension(data, filename)
            jd_text = parsed.get("content", "")
            if not jd_text.strip():
                return {"success": False, "error": "未能从文件中提取到文本内容"}

            service = MiniMaxService(api_key=api_key, resume_profile=resume_profile, model_name=model_name)
            result = await service.analyze_jd_text(jd_text)
            result["jd_text"] = jd_text
            return {"success": True, "data": result}
        except Exception as e:
            return {"success": False, "error": f"文本解析失败: {str(e)}"}

    return {"success": False, "error": f"不支持的文件格式: .{ext}"}


# ─── 动态迁移：添加 match_reason 列 ──────────────────────────────────────────
_columns_now = [col["name"] for col in sa_inspect(engine).get_columns("applications")]
if "match_reason" not in _columns_now:
    with engine.connect() as _conn:
        _conn.execute(sa_text("ALTER TABLE applications ADD COLUMN match_reason VARCHAR"))
        _conn.commit()


# ─── 批量评分 ─────────────────────────────────────────────────────────────────

class ScoreTask:
    """内存中跟踪批量评分任务进度（轻量，不入库）"""
    _tasks: dict = {}

    @classmethod
    def create(cls, task_id: str, total: int):
        cls._tasks[task_id] = {
            "total": total, "completed": 0, "failed": 0,
            "status": "processing", "errors": []
        }

    @classmethod
    def get(cls, task_id: str) -> dict | None:
        return cls._tasks.get(task_id)

    @classmethod
    def update(cls, task_id: str, completed: int, failed: int, done: bool = False):
        if task_id in cls._tasks:
            cls._tasks[task_id]["completed"] = completed
            cls._tasks[task_id]["failed"] = failed
            if done:
                cls._tasks[task_id]["status"] = "completed"

    @classmethod
    def add_error(cls, task_id: str, msg: str):
        if task_id in cls._tasks:
            cls._tasks[task_id]["errors"].append(msg)


@app.post("/api/applications/batch-score")
async def start_batch_score(
    db: Session = Depends(get_db),
    api_key: str = Depends(get_api_key),
    resume_profile: str | None = Depends(get_resume_profile),
    model_name: str | None = Depends(get_minimax_model),
    user_id: str = Depends(get_user_id),
):
    """发起批量评分任务：为当前用户所有 match_score 为 null 的记录评分"""
    pending = db.query(Application).filter(
        Application.user_id == user_id,
        Application.match_score.is_(None)
    ).all()

    if not pending:
        return {"success": True, "data": {"task_id": None, "total": 0, "message": "所有记录已有评分，无需重复评分"}}

    task_id = str(uuid.uuid4())
    ScoreTask.create(task_id, len(pending))

    # 只传 id 列表给后台，避免 session 跨线程
    app_ids = [a.id for a in pending]
    asyncio.create_task(_run_batch_score(task_id, app_ids, api_key, resume_profile, user_id, model_name))

    return {
        "success": True,
        "data": {
            "task_id": task_id,
            "total": len(pending),
            "message": f"开始为 {len(pending)} 条记录评分"
        }
    }


@app.get("/api/applications/batch-score/status/{task_id}")
async def get_batch_score_status(task_id: str, user_id: str = Depends(get_user_id)):
    """查询批量评分任务进度"""
    task = ScoreTask.get(task_id)
    if not task:
        return {"success": False, "error": "任务不存在或已过期"}
    return {"success": True, "data": task}


@app.post("/api/applications/{app_id}/score")
async def start_single_score(
    app_id: int,
    db: Session = Depends(get_db),
    api_key: str = Depends(get_api_key),
    resume_profile: str | None = Depends(get_resume_profile),
    model_name: str | None = Depends(get_minimax_model),
    user_id: str = Depends(get_user_id),
):
    """重新检测特定单条记录的匹配度"""
    record = db.query(Application).filter_by(id=app_id, user_id=user_id).first()
    if not record:
        return {"success": False, "error": "记录不存在"}
    
    task_id = str(uuid.uuid4())
    ScoreTask.create(task_id, 1)
    asyncio.create_task(_run_batch_score(task_id, [app_id], api_key, resume_profile, user_id, model_name))
    
    return {
        "success": True,
        "data": {
            "task_id": task_id,
            "total": 1,
            "message": "已开始重新检测"
        }
    }


@app.get("/api/applications/unscored-count")
async def get_unscored_count(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    """返回当前用户未评分的记录数量"""
    count = db.query(Application).filter(
        Application.user_id == user_id,
        Application.match_score.is_(None)
    ).count()
    return {"success": True, "data": {"unscored": count}}


async def _run_batch_score(task_id: str, app_ids: list[int], api_key: str, resume_profile: str | None, user_id: str, model_name: str | None = None):
    """后台异步执行批量评分"""
    service = MiniMaxService(api_key=api_key, resume_profile=resume_profile, model_name=model_name)
    db = SessionLocal()
    completed = 0
    failed = 0

    try:
        for app_id in app_ids:
            record = db.query(Application).filter_by(id=app_id, user_id=user_id).first()
            if not record:
                failed += 1
                continue

            try:
                result = await service.score_single_application(
                    company_name=record.company_name,
                    position=record.position,
                    city=record.city,
                    company_size=record.company_size,
                    salary_range=record.salary_range,
                    jd_text=record.jd_text,
                )
                record.match_score = result.get("match_score")
                # match_reason 列动态迁移后存在，用 setattr 兼容旧 model
                setattr(record, "match_reason", result.get("match_reason", ""))
                db.commit()
                completed += 1
            except Exception as e:
                failed += 1
                ScoreTask.add_error(task_id, f"ID={app_id}: {str(e)[:80]}")
                db.rollback()

            ScoreTask.update(task_id, completed, failed)
            # 每条之间短暂让步，避免 API 限流
            await asyncio.sleep(0.3)

        ScoreTask.update(task_id, completed, failed, done=True)
    finally:
        db.close()


# ─── 状态枚举接口 ─────────────────────────────────────────────────────────────

@app.get("/api/meta/statuses")
async def get_all_statuses():
    """返回全量投递进度枚举列表"""
    return {"success": True, "data": ALL_STATUSES}


# ─── 投递记录 CRUD ─────────────────────────────────────────────────────────────

@app.post("/api/applications")
async def create_application(data: ApplicationCreate, db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    """手动新建一条投递记录"""
    try:
        skills_str = json.dumps(data.skills, ensure_ascii=False) if data.skills else None
        company_tier = data.company_tier or extract_company_tier(data.company_size)
        record = Application(
            user_id=user_id,
            company_name=data.company_name,
            position=data.position,
            city=data.city,
            company_size=data.company_size,
            company_tier=company_tier,
            jd_text=data.jd_text,
            salary_range=data.salary_range,
            education=data.education,
            experience=data.experience,
            skills=skills_str,
            match_score=data.match_score,
            status=data.status or "已投递",
            channel=data.channel,
            apply_date=data.apply_date or date.today(),
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        return {"success": True, "data": {"id": record.id}}
    except Exception as e:
        return {"success": False, "error": str(e)}


@app.put("/api/applications/{app_id}")
async def update_application(app_id: int, data: dict, db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    """更新单条投递记录的任意字段"""
    record = db.query(Application).filter_by(id=app_id, user_id=user_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="记录不存在")
    updatable = {"company_name", "position", "city", "company_size", "company_tier",
                 "jd_text", "salary_range", "education", "experience", "skills",
                 "match_score", "status", "channel", "apply_date"}
    for key, val in data.items():
        if key not in updatable:
            continue
        if key == "skills" and isinstance(val, list):
            setattr(record, key, json.dumps(val, ensure_ascii=False))
        elif key == "apply_date" and isinstance(val, str):
            try:
                setattr(record, key, date.fromisoformat(val))
            except ValueError:
                pass
        elif key == "company_size":
            record.company_size = val
            record.company_tier = extract_company_tier(val)
        else:
            setattr(record, key, val)
    db.commit()
    return {"success": True}


@app.delete("/api/applications/{app_id}")
async def delete_application(app_id: int, db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    """删除单条投递记录"""
    record = db.query(Application).filter_by(id=app_id, user_id=user_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="记录不存在")
    db.delete(record)
    db.commit()
    return {"success": True}


# ─── API Key 清除 ─────────────────────────────────────────────────────────────

@app.delete("/api/settings/api-key")
async def clear_api_key_delete(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    """清除当前用户的 API Key"""
    config = db.query(UserConfig).filter_by(user_id=user_id, key="minimax_api_key").first()
    if config:
        config.value = ""
        db.commit()
    return {"success": True, "message": "API Key 已清除"}


# ─── 自定义 Prompt 存取 ───────────────────────────────────────────────────────

@app.get("/api/settings/custom-prompt")
async def get_custom_prompt(db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    config = db.query(UserConfig).filter_by(user_id=user_id, key="custom_prompt").first()
    return {"success": True, "data": {"prompt": config.value if config else ""}}


@app.post("/api/settings/custom-prompt")
async def save_custom_prompt(request: dict, db: Session = Depends(get_db), user_id: str = Depends(get_user_id)):
    prompt = request.get("prompt", "").strip()
    config = db.query(UserConfig).filter_by(user_id=user_id, key="custom_prompt").first()
    if config:
        config.value = prompt
    else:
        config = UserConfig(user_id=user_id, key="custom_prompt", value=prompt)
        db.add(config)
    db.commit()
    return {"success": True, "message": "自定义 Prompt 已保存"}


# ─── Excel 导出 ───────────────────────────────────────────────────────────────

@app.get("/api/export/xlsx")
async def export_xlsx(
    sheets: str = Query("records,city,status,tier"),  # 逗号分隔的 sheet 名称
    db: Session = Depends(get_db),
    user_id: str = Depends(get_user_id),
):
    """导出投递数据为 Excel，支持多 sheet 选择"""
    if not OPENPYXL_AVAILABLE:
        return {"success": False, "error": "服务端未安装 openpyxl，请执行: pip install openpyxl"}

    requested = {s.strip() for s in sheets.split(",")}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 删除默认空 sheet

    # 通用样式
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(ws, headers: list[str]):
        ws.append(headers)
        for col_idx, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border

    def auto_width(ws):
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)

    all_apps = db.query(Application).filter_by(user_id=user_id).order_by(Application.apply_date.desc()).all()

    # Sheet 1: 投递明细
    if "records" in requested:
        ws = wb.create_sheet("投递明细")
        headers = ["公司名称", "岗位", "城市", "公司规模", "公司档次", "投递进度", "投递渠道",
                   "匹配度", "投递日期"]
        style_header(ws, headers)
        for app in all_apps:
            ws.append([
                app.company_name, app.position, app.city or "", app.company_size or "",
                app.company_tier or "", app.status or "", app.channel or "",
                app.match_score, str(app.apply_date) if app.apply_date else "",
            ])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center")
        auto_width(ws)
        ws.freeze_panes = "A2"

    # Sheet 2: 城市分布
    if "city" in requested:
        ws = wb.create_sheet("城市分布")
        style_header(ws, ["城市", "投递数量"])
        from collections import Counter
        city_cnt = Counter(a.city or "未知" for a in all_apps)
        for city, cnt in sorted(city_cnt.items(), key=lambda x: -x[1]):
            ws.append([city, cnt])
        auto_width(ws)

    # Sheet 3: 状态分布
    if "status" in requested:
        ws = wb.create_sheet("投递进度分布")
        style_header(ws, ["投递进度", "数量"])
        from collections import Counter
        status_cnt = Counter(a.status or "未知" for a in all_apps)
        for status, cnt in sorted(status_cnt.items(), key=lambda x: -x[1]):
            ws.append([status, cnt])
        auto_width(ws)

    # Sheet 4: 公司档次分布
    if "tier" in requested:
        ws = wb.create_sheet("公司档次分布")
        style_header(ws, ["档次", "数量"])
        from collections import Counter
        tier_cnt = Counter(a.company_tier or "未知" for a in all_apps)
        tier_order = ["大厂", "中大厂", "中厂", "中小厂", "小厂", "未知"]
        for tier in tier_order:
            if tier in tier_cnt:
                ws.append([tier, tier_cnt[tier]])
        auto_width(ws)

    # Sheet 5: 匹配度分布
    if "score" in requested:
        ws = wb.create_sheet("匹配度分布")
        style_header(ws, ["区间", "数量"])
        scored = [a.match_score for a in all_apps if a.match_score is not None]
        ranges = [("0-20", 0, 20), ("21-40", 21, 40), ("41-60", 41, 60),
                  ("61-80", 61, 80), ("81-100", 81, 100)]
        for label, lo, hi in ranges:
            ws.append([label, sum(1 for s in scored if lo <= s <= hi)])
        auto_width(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"求职看板导出_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    encoded_filename = urllib.parse.quote(filename)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )


async def _auto_mark_no_reply():
    """后台任务：每小时扫描一次，将已投递超过7天未回复的记录标记为'7天未回复'"""
    while True:
        try:
            db = SessionLocal()
            try:
                seven_days_ago = date.today() - timedelta(days=7)
                updated = db.query(Application).filter(
                    Application.status == "已投递",
                    Application.apply_date <= seven_days_ago
                ).update({Application.status: "7天未回复"}, synchronize_session=False)
                db.commit()
                if updated > 0:
                    print(f"[AutoMark] {datetime.now().isoformat()} 自动标记 {updated} 条记录为 7天未回复")
            finally:
                db.close()
        except Exception as e:
            print(f"[AutoMark] 任务执行异常: {e}")

        await asyncio.sleep(3600)


@app.on_event("startup")
async def startup_event():
    asyncio.create_task(_auto_mark_no_reply())


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
