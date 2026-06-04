import io
import csv
import chardet
from datetime import datetime


def detect_encoding(data: bytes) -> str:
    result = chardet.detect(data)
    return result.get("encoding") or "utf-8"


def parse_csv(data: bytes) -> list[dict]:
    """解析 CSV/Excel 导出的 CSV 文件，返回字典列表"""
    encoding = detect_encoding(data)
    try:
        text = data.decode(encoding)
    except UnicodeDecodeError:
        text = data.decode("utf-8-sig")

    reader = csv.DictReader(io.StringIO(text))
    return list(reader)


def parse_excel(data: bytes) -> list[dict]:
    """解析 Excel 文件，返回字典列表"""
    try:
        import pandas as pd
    except ImportError:
        raise ImportError("请安装 pandas: pip install pandas openpyxl")

    df = pd.read_excel(io.BytesIO(data))
    df = df.where(pd.notnull(df), None)
    return df.to_dict(orient="records")


def parse_docx(data: bytes) -> str:
    """解析 Word 文档，返回纯文本（包含段落和表格内容）"""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("请安装 python-docx: pip install python-docx")

    doc = Document(io.BytesIO(data))
    parts = []

    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in seen:
                    seen.add(text)
                    parts.append(text)

    return "\n".join(parts)


def parse_text(data: bytes) -> str:
    """解析纯文本文件（TXT/Markdown），自动检测编码"""
    encoding = detect_encoding(data)
    try:
        return data.decode(encoding)
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="replace")


def parse_file_by_extension(data: bytes, filename: str) -> dict:
    """根据文件扩展名路由到对应的解析器

    Returns:
        {"type": "structured", "rows": [...]}  或
        {"type": "text", "content": "..."}
    """
    name = filename.lower()

    if name.endswith(".csv"):
        rows = parse_csv(data)
        return {"type": "structured", "rows": rows}

    if name.endswith(".xlsx") or name.endswith(".xls"):
        rows = parse_excel(data)
        return {"type": "structured", "rows": rows}

    if name.endswith(".docx"):
        text = parse_docx(data)
        return {"type": "text", "content": text}

    if name.endswith(".txt") or name.endswith(".md"):
        text = parse_text(data)
        return {"type": "text", "content": text}

    raise ValueError(f"不支持的文件格式: {filename}")
